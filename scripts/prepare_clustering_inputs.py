"""Prepare aggregate inputs for outcome-free clusters and mortality comparisons.
Run from supplement root with bundled Python/openpyxl:
 python scripts/prepare_clustering_inputs.py
Public Census downloads are cached; Primary files are read without modification.
"""
from pathlib import Path
from collections import Counter
import csv,io,os,json,urllib.request,zipfile,hashlib,datetime,calendar,concurrent.futures
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
PRIMARY=Path(os.environ.get("STARS_SOURCE_DIR") or ROOT.parent)
CACHE=ROOT/"data/external/census"
OUT=ROOT/"data/processed/clustering"
CACHE.mkdir(parents=True,exist_ok=True);OUT.mkdir(parents=True,exist_ok=True)
BASE="https://www2.census.gov/programs-surveys/acs/summary_file/2011/"
manifest={}
def fetch(name,url):
    target=CACHE/name
    if not target.exists():
        data=urllib.request.urlopen(url,timeout=45).read()
        target.write_bytes(data)
    manifest[name]={"file":str(target.relative_to(ROOT)),"url":url,
        "sha256":hashlib.sha256(target.read_bytes()).hexdigest()}
    return target
def readcsv(f,encoding="utf-8-sig"):
    with open(f,newline="",encoding=encoding) as h:return list(csv.DictReader(h))
def write(name,rows):
    with (OUT/name).open("w",newline="",encoding="utf-8") as h:
        w=csv.DictWriter(h,fieldnames=rows[0]);w.writeheader();w.writerows(rows)
def county(s):
    return s.strip().lstrip(".").replace(" County, Montana","").replace("Lewis & Clark","Lewis and Clark")
lookup=fetch("Sequence_Number_and_Table_Number_Lookup.txt",
    BASE+"documentation/5_year/user_tools/Sequence_Number_and_Table_Number_Lookup.txt")
rows=list(csv.DictReader(io.StringIO(lookup.read_bytes().decode("cp1252").replace("\r\r\n","\n"))))
tables=["B01002","B02001","B19013","B17001","B11001","B12001","B15002"]
starts={r["Table ID"]:r for r in rows if r["Table ID"] in tables and r["position"]!="."}
assert set(starts)==set(tables)
geo=fetch("g20115mt.csv",BASE+"data/5_year_seq_by_state/Montana/All_Geographies_Not_Tracts_Block_Groups/g20115mt.csv")
geos=list(csv.reader(io.StringIO(geo.read_bytes().decode("cp1252"))))
# Geographic file: SUMLEVEL=050, COMPONENT=00, LOGRECNO column 5.
# Identify county names from the standard GEOID/name columns to verify parsing.
geos=[r for r in geos if r[2]=="050" and r[3]=="00"]
assert len(geos)==56
name_index=next(i for i,x in enumerate(geos[0]) if x.endswith(" County, Montana"))
id_index=next(i for i,x in enumerate(geos[0]) if x.startswith("05000US30"))
by_log={r[4]:(r[id_index][-3:],county(r[name_index])) for r in geos}
seqs=sorted(set(r["seq"] for r in starts.values()))
def getseq(seq):
    name="20115mt"+seq+"000.zip"
    f=fetch(name,BASE+"data/5_year_seq_by_state/Montana/All_Geographies_Not_Tracts_Block_Groups/"+name)
    with zipfile.ZipFile(f) as z:
        e=next(n for n in z.namelist() if Path(n).name.startswith("e"))
        m=next(n for n in z.namelist() if Path(n).name.startswith("m"))
        est={r[5]:r for r in csv.reader(io.StringIO(z.read(e).decode("cp1252"))) if r[5] in by_log}
        moe={r[5]:r for r in csv.reader(io.StringIO(z.read(m).decode("cp1252"))) if r[5] in by_log}
    assert len(est)==56
    return seq,est,moe
with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
    sequences={seq:(e,m) for seq,e,m in pool.map(getseq,seqs)}
def value(log,table,line,moe=False):
    r=starts[table];index=int(r["position"])-1+line-1
    x=sequences[r["seq"]][int(moe)][log][index]
    val=float(x)
    assert val>=0,(log,table,line,x)
    return val
# Add all-county uninsured percentage from existing early-period ACS workbook.
insurance_file=PRIMARY/"CensusCountyData/ACSST5Y2012_HealthInsuranceCoverage_County.xlsx"
wb=load_workbook(insurance_file,read_only=True,data_only=True)
rr=list(wb["DataValues"].values);header=rr[0]
insurance={str(r[header.index("County")]).strip():100*float(r[header.index("PercentUninsured")])
           for r in rr[1:] if r[header.index("County")]}
wb.close();assert len(insurance)==56
context=[]
for log,(fips,name) in sorted(by_log.items(),key=lambda x:x[1][0]):
    v=lambda t,n:value(log,t,n)
    context.append(dict(FIPS=fips,COUNTY=name,MEDAGE=v("B01002",1),
        AIAN_ALONE_PCT=100*v("B02001",4)/v("B02001",1),
        MEDINCOME=v("B19013",1),POVERTY_PCT=100*v("B17001",2)/v("B17001",1),
        SINGLE_HH_PCT=100*v("B11001",8)/v("B11001",1),
        DIVSEP_PCT=100*sum(v("B12001",i) for i in [7,10,16,19])/v("B12001",1),
        DEGREE_PCT=100*sum(v("B15002",i) for i in list(range(14,19))+list(range(31,36)))/v("B15002",1),
        UNINSURED_PCT=insurance[name]))
write("county_context.csv",context)
# Preserve the source estimates and 90%-MOE values used for audit of small counties.
aud=[]
lines={"B01002":[1],"B02001":[1,4],"B19013":[1],"B17001":[1,2],
       "B11001":[1,8],"B12001":[1,7,10,16,19],
       "B15002":[1]+list(range(14,19))+list(range(31,36))}
for log,(fips,name) in by_log.items():
    for table,nums in lines.items():
        for n in nums:
            aud.append(dict(FIPS=fips,COUNTY=name,table=table,line=n,
                estimate=value(log,table,n),moe90=value(log,table,n,True)))
write("county_acs_source_cells.csv",aud)
# Revised intercensal series for 2012-2019; vintage 2021 for 2020-2021.
inter=fetch("co-est2020int-pop-30.xlsx",
 "https://www2.census.gov/programs-surveys/popest/tables/2010-2020/intercensal/county/co-est2020int-pop-30.xlsx")
wb=load_workbook(inter,read_only=True,data_only=True)
rr=list(wb.active.values);year_columns={int(y):i for i,y in enumerate(rr[3]) if isinstance(y,(int,float))}
pop={}
for r in rr:
    if isinstance(r[0],str) and r[0].startswith(".") and r[0].endswith(" County, Montana"):
        for y in range(2012,2020):pop[(county(r[0]),y)]=int(r[year_columns[y]])
wb.close()
v21=fetch("co-est2021-alldata.csv",
 "https://www2.census.gov/programs-surveys/popest/datasets/2020-2021/counties/totals/co-est2021-alldata.csv")
for r in readcsv(v21,"cp1252"):
    if int(r["STATE"])==30 and int(r["COUNTY"])>0 and int(r["SUMLEV"])==50:
        for y in [2020,2021]:pop[(r["CTYNAME"].replace(" County",""),y)]=int(r["POPESTIMATE"+str(y)])
assert len(pop)==560 and all(x>0 for x in pop.values())
source=PRIMARY/"R - Data Exploration/20230831_Montana_Suicide_Database.csv"
stars=readcsv(source)
counties={r["COUNTY"] for r in context}
valid=[r for r in stars if county(r["DEM_RES03CountyCorrected"]) in counties]
unknown=[r for r in stars if r not in valid]
counts=Counter((county(r["DEM_RES03CountyCorrected"]),int(r["DEATH_2Year"])) for r in valid)
write("county_year_mortality.csv",[dict(FIPS=c["FIPS"],COUNTY=c["COUNTY"],YEAR=y,
    DEATHS=counts[c["COUNTY"],y],POPULATION=pop[c["COUNTY"],y])
    for c in context for y in range(2012,2022)])
write("residence_accounting.csv",[dict(category="Montana county matched",n=len(valid)),
    dict(category="County unknown or outside Montana",n=len(unknown))])
# Date completeness determines eligible years BEFORE temporal outcome tests.
def known(x):return x not in ("","NA",None)
coverage=[]
for y in range(2012,2022):
    z=[r for r in stars if int(r["DEATH_2Year"])==y]
    coverage.append(dict(YEAR=y,n=len(z),month_missing=sum(not known(r["DEATH_4MonthNum"]) for r in z),
        weekday_missing=sum(not known(r["DEATH_5Day"]) for r in z)))
write("date_coverage.csv",coverage)
month_years=[r["YEAR"] for r in coverage if r["month_missing"]==0]
weekday_years=[r["YEAR"] for r in coverage if r["weekday_missing"]==0]
assert month_years==[2018,2020,2021]
assert weekday_years==[2013,2014,2015,2016,2018,2020,2021]
# Temporal mortality counts are restricted to Montana residents, for matching denominators.
# Unknown/out-of-state residence is not included in statewide population rates.
def temporal(years,axis):
    out=[]
    for y in years:
        z=[r for r in valid if int(r["DEATH_2Year"])==y]
        annual_pop=sum(pop[(c,y)] for c in counties)
        nday=366 if calendar.isleap(y) else 365
        dates=[datetime.date(y,1,1)+datetime.timedelta(days=i) for i in range(nday)]
        for unit in (range(1,13) if axis=="month" else range(7)):
            exposure=sum((d.month==unit if axis=="month" else d.weekday()==unit) for d in dates)
            count=sum((int(r["DEATH_4MonthNum"])==unit if axis=="month"
                else r["DEATH_5Day"].strip().lower()==calendar.day_name[unit].lower())
                for r in z if known(r["DEATH_4MonthNum"] if axis=="month" else r["DEATH_5Day"]))
            out.append(dict(YEAR=y,UNIT=unit if axis=="month" else calendar.day_name[unit],
                DEATHS=count,DAYS=exposure,PERSON_YEARS=annual_pop*exposure/nday))
    return out
write("month_year_mortality.csv",temporal(month_years,"month"))
write("weekday_year_mortality.csv",temporal(weekday_years,"weekday"))
# Temporal context is exported from the supplied workbooks.
import runpy
runpy.run_path(str(ROOT/"scripts/prepare_temporal_context.py"),run_name="__main__")
# Source documentation and hashes. Raw records and dates are never exported.
manifest["insurance_workbook"]={"file":str(insurance_file.relative_to(PRIMARY)),
    "url":"Census workbook, ACS 2008-2012",
    "sha256":hashlib.sha256(insurance_file.read_bytes()).hexdigest()}
manifest["STARS_CSV"]={"file":str(source.relative_to(PRIMARY)),"url":"private primary source",
    "sha256":hashlib.sha256(source.read_bytes()).hexdigest()}
for name in ["month_context.csv","weekday_context.csv","temporal_context_sources.csv"]:
    f=OUT/name
    manifest[name]={"file":str(f.relative_to(ROOT)),"url":"external contextual measures; see docs/INPUTS.md",
      "sha256":hashlib.sha256(f.read_bytes()).hexdigest()}
write("source_manifest.csv",[manifest[k] for k in sorted(manifest)])
description=[
 ("MEDAGE","2007-2011 ACS B01002_001","Median age, years"),
 ("AIAN_ALONE_PCT","B02001_004 / B02001_001","American Indian/Alaska Native alone, percent"),
 ("MEDINCOME","B19013_001","Median household income, 2011 dollars"),
 ("POVERTY_PCT","B17001_002 / B17001_001","Below poverty, percent of universe"),
 ("SINGLE_HH_PCT","B11001_008 / B11001_001","One-person households, percent"),
 ("DIVSEP_PCT","B12001 (007+010+016+019)/001","Divorced/separated, percent aged 15+"),
 ("DEGREE_PCT","B15002 (014:018 + 031:035)/001","Associate degree or higher, percent aged 25+"),
 ("UNINSURED_PCT","ACSST5Y2012 insurance workbook","Uninsured, percent civilian noninstitutionalized")]
write("county_feature_definitions.csv",[dict(feature=a,source=b,definition=c) for a,b,c in description])
print("Prepared",len(context),"counties;",len(valid),"county-matched deaths;",
      len(unknown),"excluded residence records; month years",month_years,"weekday years",weekday_years)
print("Petroleum:",next(r for r in context if r["COUNTY"]=="Petroleum"))
